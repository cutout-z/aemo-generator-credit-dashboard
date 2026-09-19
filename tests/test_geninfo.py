"""Tests for the AEMO Generation Information quarterly fetcher + commitment diff.

Hermetic: synthetic workbooks / frames + tmp dirs only — no network, no repo
writes. The HTTP layer is monkeypatched at the module-function boundary
(``geninfo.discover_edition`` / ``fetch_xlsx`` / ``probe_edition_url``) so the
real orchestration, storage, diff and publish code paths run.
"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from src import geninfo
from src.geninfo import (
    GenInfoError,
    attach_geninfo_doc,
    build_gen_info_payload,
    candidate_edition_urls,
    diff_commitment_status,
    discover_previous_edition,
    events_by_duid,
    extract_edition_links,
    load_snapshot,
    normalize_geninfo,
    parse_geninfo_workbook,
    publish_gen_info_json,
    run_geninfo_lane,
    select_current_edition,
    snapshot_fingerprint,
)
from src.run_status import LaneRun, check_factor_block_continuity, STATUS_DEGRADED, STATUS_ERROR, STATUS_OK

# Header names as they appear in the real workbook (some carry a leading
# space — parse must strip). Header lives in row 4 of the 'Generator
# Information' sheet.
_WB_COLUMNS = [
    " Survey ID",
    "Site Name",
    " KCI Id",
    "Site Owner",
    "Custodian",
    "Region",
    "Max Site Capacity (AC)",
    "Gen Info Unit ID",
    "Unit Name",
    "Technology Type",
    "Technology Detail",
    "DUID",
    "Dispatch Type",
    "Unit Count",
    "Unit Capacity (MW AC)",
    "Aggregated Nameplate Capacity (MW AC)",
    "Agg Nameplate Storage Capacity (MWh)",
    "Commitment Status",
    "Full Commercial Use Date",
    "Expected Closure Year",
    "Closure Date",
    "Survey Latest Update Date",
]

_DEFAULTS = {
    " Survey ID": "9001",
    " KCI Id": "N00001",
    "Site Owner": "Test Energy Pty Ltd",
    "Custodian": "TESTCO",
    "Region": "NSW1",
    "Max Site Capacity (AC)": 100.0,
    "Technology Type": "Battery Storage",
    "Technology Detail": "Lithium-ion",
    "DUID": None,
    "Dispatch Type": "Scheduled",
    "Unit Count": 1,
    "Unit Capacity (MW AC)": 100.0,
    "Aggregated Nameplate Capacity (MW AC)": 100.0,
    "Agg Nameplate Storage Capacity (MWh)": 200.0,
    "Commitment Status": "Publicly Announced",
    "Full Commercial Use Date": None,
    "Expected Closure Year": 2060,
    "Closure Date": None,
    "Survey Latest Update Date": datetime(2026, 5, 22, 13, 38, 49),
}


def _wb_row(unit_id: str, unit_name: str, **overrides) -> dict:
    row = {**{k: None for k in _WB_COLUMNS}, **_DEFAULTS}
    row["Gen Info Unit ID"] = unit_id
    row["Unit Name"] = unit_name
    row["Site Name"] = overrides.pop("Site Name", f"Site {unit_id}")
    row.update(overrides)
    return row


def _write_workbook(path: Path, rows: list[dict], sheet: str = "Generator Information",
                    columns: list[str] | None = None) -> Path:
    cols = columns if columns is not None else _WB_COLUMNS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(row=1, column=1, value="Existing Generators & New Developments")
    ws.cell(row=2, column=1, value="Since the October 2025 release the workbook has been restructured.")
    ws.cell(row=3, column=1, value="For seasonal capacity trends see the Summary sheet.")
    for j, col in enumerate(cols, start=1):
        ws.cell(row=4, column=j, value=col)
    for i, row in enumerate(rows, start=5):
        for j, col in enumerate(cols, start=1):
            ws.cell(row=i, column=j, value=row.get(col))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _frame(rows: list[dict], edition: str = "2026-07") -> pd.DataFrame:
    """Normalised snapshot straight from workbook-shaped rows."""
    return normalize_geninfo(pd.DataFrame(rows), edition)


# ────────────────────────────────────────────────────────────────────────
# Edition links / discovery
# ────────────────────────────────────────────────────────────────────────

_LANDING_HTML = """
<html><body>
<a href="https://www.aemo.com.au/-/media/files/electricity/nem/planning_and_forecasting/generation_information/2026/nem-generation-information-july-2026.xlsx?rev=3455851f2bc945b7ab61c5ceed272992&amp;sc_lang=en">July 2026</a>
<a href="/-/media/files/electricity/nem/planning_and_forecasting/generation_information/2026/nem-generation-information-apr-2026.xlsx?rev=6e58a954c66d4d0db93856404fa3e7b5&sc_lang=en">Apr 2026</a>
<a href="https://www.aemo.com.au/-/media/files/electricity/nem/planning_and_forecasting/generation_information/2026/nem-generation-information-jan-2026.xlsx?rev=66d8f10c76c24d8688cdb7a2302dd7d2&sc_lang=en">Jan 2026</a>
<a href="https://www.aemo.com.au/-/media/files/electricity/nem/planning_and_forecasting/generation_information/2025/nem-generation-information-oct-2025.xlsx?rev=38f56f5fa2aa4ceb93272ba188e85852&sc_lang=en">Oct 2025</a>
<a href="https://www.aemo.com.au/-/media/files/electricity/nem/planning_and_forecasting/generation_information/2025/nem-generation-information-april-2025.xlsx?rev=21b44f9af0114c17b440c80a5f21007c&sc_lang=en">April 2025</a>
<a href="https://www.aemo.com.au/-/media/files/electricity/nem/planning_and_forecasting/generation_information/2024/nem-generation-information-7-feb-2024.xlsx?rev=0e5518e295ee461d86ac7e84a3654624&sc_lang=en">Feb 2024</a>
<a href="https://www.aemo.com.au/-/media/files/electricity/nem/planning_and_forecasting/generation_information/2026/generating-unit-expected-closure-year-july-2026.xlsx?rev=86493e47cf104e6b80002f921410ef1c&sc_lang=en">Closure year</a>
<a href="https://www.aemo.com.au/-/media/files/electricity/nem/planning_and_forecasting/generation_information/2026/kci-datafile-compiled-nem.xlsx?rev=2c020fb4f0d849ae9cfbd76bb0c670e7&sc_lang=en">KCI</a>
<a href="https://www.aemo.com.au/-/media/files/electricity/nem/planning_and_forecasting/generation_information/may-2019/generation_information_nsw_may_2019.xlsx?rev=e86917de051a451798db76af9d617179&sc_lang=en">Legacy NSW</a>
</body></html>
"""


class TestEditionLinks:
    def test_extract_filters_to_national_series_and_orders(self):
        links = extract_edition_links(_LANDING_HTML)
        editions = [e["edition"] for e in links]
        assert editions == [
            "2024-02", "2025-04", "2025-10", "2026-01", "2026-04", "2026-07",
        ]
        # companion workbooks and the pre-2025 per-region files are not editions
        assert all("closure" not in e["filename"] for e in links)
        assert all("kci" not in e["filename"] for e in links)
        # href query (the revision hash) is preserved, entities decoded
        july = links[-1]
        assert "rev=3455851f2bc945b7ab61c5ceed272992" in july["url"]
        assert "sc_lang=en" in july["url"]

    def test_relative_hrefs_resolved(self):
        links = extract_edition_links(_LANDING_HTML)
        apr = [e for e in links if e["edition"] == "2026-04"][0]
        assert apr["url"].startswith("https://www.aemo.com.au/-/media/")

    def test_select_current_edition_never_future(self):
        links = extract_edition_links(_LANDING_HTML)
        assert select_current_edition(links, today=date(2026, 9, 19))["edition"] == "2026-07"
        assert select_current_edition(links, today=date(2026, 3, 1))["edition"] == "2026-01"
        assert select_current_edition([], today=date(2026, 9, 19)) is None

    def test_candidate_urls_month_spellings(self):
        assert candidate_edition_urls(2026, 4) == [
            geninfo.GENINFO_MEDIA_BASE + "/2026/nem-generation-information-apr-2026.xlsx",
            geninfo.GENINFO_MEDIA_BASE + "/2026/nem-generation-information-april-2026.xlsx",
        ]
        assert candidate_edition_urls(2026, 7)[0].endswith("nem-generation-information-jul-2026.xlsx")

    def test_recent_quarter_months(self):
        months = geninfo._recent_quarter_months(date(2026, 9, 19), count=4)
        assert months == [(2026, 7), (2026, 4), (2026, 1), (2025, 10)]

    def test_parse_edition_from_filename_variants(self):
        parse = geninfo._parse_edition_from_filename
        assert parse("nem-generation-information-july-2026.xlsx") == (2026, 7)
        assert parse("nem-generation-information-october-2024.xlsx") == (2024, 10)
        assert parse("nem-generation-information-7-feb-2024.xlsx") == (2024, 2)
        assert parse("generation_information_nsw_may_2019.xlsx") is None
        assert parse("nem-generation-information-20200729.xlsx") is None


# ────────────────────────────────────────────────────────────────────────
# Parse + normalise
# ────────────────────────────────────────────────────────────────────────

class TestParseNormalize:
    def test_parse_header_row_and_values(self, tmp_path):
        path = _write_workbook(tmp_path / "edition.xlsx", [
            _wb_row("248901", "AMBS_GU1", **{
                "Region": "QLD1", "DUID": "BRDDBES1",
                "Commitment Status": "Committed",
                "Full Commercial Use Date": datetime(2027, 7, 31),
            }),
            _wb_row("248902", "AMBS_GU2", **{"Commitment Status": "Anticipated"}),
        ])
        raw = parse_geninfo_workbook(path)
        assert "Gen Info Unit ID" in raw.columns  # header row 4 found, names stripped
        assert [str(v) for v in raw["Gen Info Unit ID"]] == ["248901", "248902"]

        snap = normalize_geninfo(raw, "2026-07")
        assert snap["edition"].iloc[0] == "2026-07"
        first = snap.iloc[0]
        assert first["unit_id"] == "248901"
        assert first["duid"] == "BRDDBES1"
        assert first["commitment_status"] == "Committed"
        assert first["capacity_mw"] == pytest.approx(100.0)
        assert first["storage_mwh"] == pytest.approx(200.0)
        assert first["full_commercial_use_date"] == "2027-07-31"
        assert first["expected_closure_year"] == 2060
        # blank DUID (62% of real rows) stays absent, never a stub string
        assert snap.iloc[1]["duid"] is None

    def test_duplicate_unit_id_rows_are_kept(self, tmp_path):
        """Two physical units may share one Gen Info Unit ID (Whitwood case)."""
        path = _write_workbook(tmp_path / "edition.xlsx", [
            _wb_row("122903", "GM2", **{"Aggregated Nameplate Capacity (MW AC)": 1.1}),
            _wb_row("122903", "GM3", **{"Aggregated Nameplate Capacity (MW AC)": 0.7}),
        ])
        snap = normalize_geninfo(parse_geninfo_workbook(path), "2026-07")
        assert len(snap) == 2
        assert sorted(snap["unit_name"]) == ["GM2", "GM3"]

    def test_sheet_fallback_and_missing_schema(self, tmp_path):
        path = _write_workbook(tmp_path / "renamed.xlsx", [_wb_row("1", "U1")],
                               sheet="Register")
        snap = normalize_geninfo(parse_geninfo_workbook(path), "2026-07")
        assert len(snap) == 1

        empty = tmp_path / "no_schema.xlsx"
        wb = openpyxl.Workbook()
        wb.active.cell(row=4, column=1, value="Nothing")
        wb.save(empty)
        with pytest.raises(GenInfoError):
            parse_geninfo_workbook(empty)

    def test_normalize_rejects_missing_columns(self):
        with pytest.raises(GenInfoError):
            normalize_geninfo(pd.DataFrame({"Site Name": ["x"]}), "2026-07")
        with pytest.raises(GenInfoError):
            normalize_geninfo(pd.DataFrame(), "2026-07")

    def test_edition_column_aliases_resolve(self, tmp_path):
        """Apr-2026-era headers (renamed vs Jul-2026) still normalise."""
        renames = {
            " Survey ID": "AEMO Survey ID",
            " KCI Id": "AEMO KCI ID",
            "Aggregated Nameplate Capacity (MW AC)": "Agg Nameplate Capacity (MW AC)",
        }
        cols = [renames.get(c, c) for c in _WB_COLUMNS]
        rows = [
            {renames.get(k, k): v for k, v in
             _wb_row("77", "ALIAS", **{"Commitment Status": "Committed"}).items()}
        ]
        path = _write_workbook(tmp_path / "apr.xlsx", rows, columns=cols)
        snap = normalize_geninfo(parse_geninfo_workbook(path), "2026-04")
        row = snap.iloc[0]
        assert row["capacity_mw"] == pytest.approx(100.0)
        assert row["survey_id"] == "9001"
        assert row["kci_id"] == "N00001"

    def test_missing_optional_field_nulls_not_raises(self, tmp_path):
        cols = [c for c in _WB_COLUMNS if c != "Aggregated Nameplate Capacity (MW AC)"]
        path = _write_workbook(tmp_path / "thin.xlsx", [_wb_row("1", "U1")], columns=cols)
        snap = normalize_geninfo(parse_geninfo_workbook(path), "2026-07")
        assert snap.iloc[0]["capacity_mw"] is None
        assert snap.iloc[0]["commitment_status"] == "Publicly Announced"

    def test_fingerprint_stable_and_content_sensitive(self, tmp_path):
        path = _write_workbook(tmp_path / "edition.xlsx", [_wb_row("1", "U1")])
        snap = normalize_geninfo(parse_geninfo_workbook(path), "2026-07")
        same = normalize_geninfo(parse_geninfo_workbook(path), "2026-07")
        assert snapshot_fingerprint(snap) == snapshot_fingerprint(same)
        changed = normalize_geninfo(pd.DataFrame([
            _wb_row("1", "U1", **{"Commitment Status": "Committed"}),
        ]), "2026-07")
        assert snapshot_fingerprint(snap) != snapshot_fingerprint(changed)
        assert snapshot_fingerprint(pd.DataFrame()) == ""


# ────────────────────────────────────────────────────────────────────────
# Commitment-status diff
# ────────────────────────────────────────────────────────────────────────

def _prev_row(unit_id, unit_name, status, **overrides):
    return _wb_row(unit_id, unit_name, **{"Commitment Status": status, **overrides})


class TestCommitmentDiff:
    def test_transition_into_committed_family(self):
        prev = _frame([_prev_row("1", "A", "Anticipated")])
        curr = _frame([_prev_row("1", "A", "Committed")], edition="2026-10")
        diff = diff_commitment_status(prev, curr)
        assert diff["counts"]["new_commitments"] == 1
        assert diff["counts"]["status_changes"] == 0  # consumed by the specific bucket
        event = diff["new_commitments"][0]
        assert (event["from_status"], event["to_status"]) == ("Anticipated", "Committed")
        assert diff["compared_to_edition"] == "2026-07"

    def test_committed_to_in_service_is_commissioned_not_decommitment(self):
        prev = _frame([_prev_row("1", "A", "In Commissioning")])
        curr = _frame([_prev_row("1", "A", "In Service")], edition="2026-10")
        diff = diff_commitment_status(prev, curr)
        assert diff["counts"]["commissioned"] == 1
        assert diff["counts"]["decommitments"] == 0
        assert diff["counts"]["new_commitments"] == 0

    def test_committed_back_to_market_is_decommitment(self):
        prev = _frame([_prev_row("1", "A", "Committed")])
        curr = _frame([_prev_row("1", "A", "Anticipated")], edition="2026-10")
        diff = diff_commitment_status(prev, curr)
        assert diff["counts"]["decommitments"] == 1
        assert diff["decommitments"][0]["from_status"] == "Committed"

    def test_withdrawal_event(self):
        prev = _frame([_prev_row("1", "A", "In Service"), _prev_row("2", "B", "Committed")])
        curr = _frame([
            _prev_row("1", "A", "Announced Withdrawal"),
            _prev_row("2", "B", "Withdrawn"),
        ], edition="2026-10")
        diff = diff_commitment_status(prev, curr)
        assert diff["counts"]["withdrawals"] == 2
        assert diff["counts"]["decommitments"] == 0  # withdrawals win over de-commitment

    def test_new_unit_already_committed_and_removed_units(self):
        prev = _frame([_prev_row("9", "GONE", "Committed")])
        curr = _frame([
            _prev_row("1", "A", "In Commissioning"),
            _prev_row("2", "B", "In Service"),
        ], edition="2026-10")
        diff = diff_commitment_status(prev, curr)
        assert diff["counts"]["new_commitments"] == 1
        assert diff["new_commitments"][0]["from_status"] is None
        assert diff["counts"]["new_units"] == 2
        assert diff["counts"]["removed_units"] == 1
        assert diff["removed_units"][0]["unit_name"] == "GONE"
        assert diff["removed_units"][0]["to_status"] is None

    def test_duplicate_unit_id_matched_by_composite_key(self):
        prev = _frame([
            _prev_row("122903", "GM2", "Publicly Announced"),
            _prev_row("122903", "GM3", "Publicly Announced"),
        ])
        curr = _frame([
            _prev_row("122903", "GM2", "Committed"),
            _prev_row("122903", "GM3", "Publicly Announced"),
        ], edition="2026-10")
        diff = diff_commitment_status(prev, curr)
        assert diff["counts"]["matched_units"] == 2
        assert diff["counts"]["new_units"] == 0
        assert diff["counts"]["removed_units"] == 0
        assert diff["counts"]["new_commitments"] == 1

    def test_rename_matched_by_unit_id(self):
        prev = _frame([_prev_row("5", "OLDNAME", "Committed")])
        curr = _frame([_prev_row("5", "NEWNAME", "Committed")], edition="2026-10")
        diff = diff_commitment_status(prev, curr)
        assert diff["counts"]["new_units"] == 0
        assert diff["counts"]["removed_units"] == 0
        assert diff["counts"]["unit_name_changes"] == 1
        assert diff["unit_name_changes"][0]["from_name"] == "OLDNAME"

    def test_baseline_diff_without_previous_snapshot(self):
        curr = _frame([_prev_row("1", "A", "Committed")])
        for prev in (None, pd.DataFrame()):
            diff = diff_commitment_status(prev, curr)
            assert diff["baseline"] is True
            assert diff["compared_to_edition"] is None
            assert all(
                diff["counts"][k] == 0
                for k in ("new_commitments", "withdrawals", "decommitments",
                          "commissioned", "new_units", "removed_units")
            )

    def test_empty_current_raises(self):
        with pytest.raises(GenInfoError):
            diff_commitment_status(pd.DataFrame(), pd.DataFrame())


class TestEventsByDuid:
    def test_grouped_and_most_specific_wins(self):
        diff = {
            "edition": "2026-10",
            "new_commitments": [{"unit_id": "1", "unit_name": "A", "duid": "DUID1",
                                 "from_status": "Anticipated", "to_status": "Committed"}],
            "withdrawals": [],
            "decommitments": [],
            "commissioned": [],
            "status_changes": [
                {"unit_id": "1", "unit_name": "A", "duid": "DUID1",
                 "from_status": "Anticipated", "to_status": "Committed"},
                {"unit_id": "2", "unit_name": "B", "duid": "DUID2",
                 "from_status": "In Service", "to_status": "Anticipated"},
            ],
        }
        out = events_by_duid(diff)
        assert set(out) == {"DUID1", "DUID2"}
        assert out["DUID1"][0]["type"] == "new_commitment"  # specific beats generic
        assert len(out["DUID1"]) == 1
        assert out["DUID2"][0]["type"] == "status_change"
        assert events_by_duid(None) == {}


# ────────────────────────────────────────────────────────────────────────
# Per-generator doc block
# ────────────────────────────────────────────────────────────────────────

class TestAttachDoc:
    def test_attach_block_with_event(self, tmp_path):
        path = _write_workbook(tmp_path / "e.xlsx", [
            _wb_row("248901", "AMBS_GU1", **{"DUID": "BRDDBES1", "Commitment Status": "Committed"}),
        ])
        rows = normalize_geninfo(parse_geninfo_workbook(path), "2026-07")
        doc = {"duid": "BRDDBES1"}
        events = {"BRDDBES1": [{"type": "new_commitment", "from_status": "Anticipated",
                                "to_status": "Committed", "edition": "2026-07"}]}
        attach_geninfo_doc(doc, rows, events)
        block = doc["gen_info"]
        assert block["scope"] == "aemo_generation_information"
        assert block["edition"] == "2026-07"
        assert block["units"][0]["commitment_status"] == "Committed"
        assert block["units"][0]["event"]["type"] == "new_commitment"

    def test_attach_noop_without_rows(self):
        doc = {"duid": "X"}
        attach_geninfo_doc(doc, None)
        attach_geninfo_doc(doc, pd.DataFrame())
        assert "gen_info" not in doc

    def test_retained_stale_stamp_through_generator_json(self, tmp_path):
        from src.generate_json import generate_generator_json

        path = _write_workbook(tmp_path / "e.xlsx", [
            _wb_row("248901", "AMBS_GU1", **{"DUID": "BRDDBES1"}),
        ])
        rows = normalize_geninfo(parse_geninfo_workbook(path), "2026-07")
        out = generate_generator_json(
            "BRDDBES1", {"station_name": "Test", "region": "QLD1"}, None, None, None,
            gen_info_rows=rows,
            factor_source_status={"geninfo": {"status": "degraded", "retained": True}},
            output_dir=str(tmp_path / "generators"),
        )
        doc = json.loads(Path(out).read_text())
        assert doc["gen_info"]["source_status"] == "retained_stale"
        assert doc["gen_info"]["units"][0]["site_name"] == "Site 248901"


# ────────────────────────────────────────────────────────────────────────
# Artifact payload + publish gate
# ────────────────────────────────────────────────────────────────────────

class TestPayload:
    def _snapshot(self):
        return _frame([
            _prev_row("1", "A", "Committed", **{"Region": "NSW1", "DUID": "A1",
                                                "Aggregated Nameplate Capacity (MW AC)": 200.0,
                                                "Agg Nameplate Storage Capacity (MWh)": 400.0,
                                                "Technology Type": "Battery Storage"}),
            _prev_row("2", "B", "In Service", **{"Region": "NSW1",
                                                 "Aggregated Nameplate Capacity (MW AC)": 660.0,
                                                 "Technology Type": "Fossil"}),
            _prev_row("3", "C", "Committed*", **{"Region": "VIC1",
                                                 "Aggregated Nameplate Capacity (MW AC)": 100.0,
                                                 "Technology Type": "Wind"}),
        ])

    def test_rollups_and_json_safety(self):
        snap = self._snapshot()
        diff = diff_commitment_status(None, snap)
        payload = build_gen_info_payload(snap, {"edition": "2026-07", "source_url": "u"}, diff)
        assert payload["rows"] == 3
        assert payload["counts_by_status"]["Committed"] == 1
        assert payload["regions"]["NSW1"]["committed_mw"] == pytest.approx(200.0)
        assert payload["regions"]["NSW1"]["committed_by_technology"] == {"Battery Storage": 200.0}
        assert payload["regions"]["VIC1"]["committed_units"] == 1
        assert payload["committed_statuses"] == ["Committed", "Committed*", "In Commissioning"]
        assert len(payload["units"]) == 3
        json.dumps(payload)  # must be valid JSON (no NaN)

    def test_publish_semantic_gate(self, tmp_path):
        snap = self._snapshot()
        diff = diff_commitment_status(None, snap)
        payload = build_gen_info_payload(snap, {"edition": "2026-07"}, diff)
        out = publish_gen_info_json(payload, tmp_path)
        first = out.read_text()

        payload["updated_utc"] = "2099-01-01T00:00:00Z"  # same facts, later clock
        publish_gen_info_json(payload, tmp_path)
        assert out.read_text() == first  # stamp not bumped

        payload["units"][0]["commitment_status"] = "In Service"  # facts changed
        publish_gen_info_json(payload, tmp_path)
        assert out.read_text() != first


# ────────────────────────────────────────────────────────────────────────
# Lane orchestration (network monkeypatched)
# ────────────────────────────────────────────────────────────────────────

_JUL_URL = (
    geninfo.GENINFO_MEDIA_BASE + "/2026/nem-generation-information-july-2026.xlsx"
    "?rev=3455851f2bc945b7ab61c5ceed272992&sc_lang=en"
)
_APR_URL = (
    geninfo.GENINFO_MEDIA_BASE + "/2026/nem-generation-information-apr-2026.xlsx"
    "?rev=6e58a954c66d4d0db93856404fa3e7b5&sc_lang=en"
)


def _current_rows():
    return [
        _prev_row("1", "NEWBESS", "Committed", **{
            "DUID": "NEWB1", "Region": "SA1",
            "Aggregated Nameplate Capacity (MW AC)": 250.0}),
        _prev_row("2", "OLDGAS", "In Service", **{"DUID": "OLDG1"}),
    ]


def _baseline_rows():
    return [
        _prev_row("1", "NEWBESS", "Publicly Announced", **{"Region": "SA1"}),
        _prev_row("2", "OLDGAS", "In Service", **{"DUID": "OLDG1"}),
        _prev_row("3", "CLOSED", "Committed", **{"Region": "NSW1"}),
    ]


def _install_fake_source(monkeypatch, tmp_path, *, edition="2026-07",
                         url=_JUL_URL, prev_edition="2026-04", prev_url=_APR_URL,
                         fail_prev=False, last_modified=None):
    """Fake the network: discovery + xlsx download write fixture workbooks."""
    calls = {"fetch": 0}

    monkeypatch.setattr(geninfo, "discover_edition", lambda today=None: {
        "url": url, "edition": edition, "year": int(edition[:4]),
        "month": int(edition[5:]), "filename": url.split("/")[-1].split("?")[0],
        "route": "landing_page",
    })
    if prev_edition is None:
        monkeypatch.setattr(geninfo, "discover_previous_edition",
                            lambda current, today=None: None)
    else:
        monkeypatch.setattr(geninfo, "discover_previous_edition",
                            lambda current, today=None: {
                                "url": prev_url, "edition": prev_edition,
                                "year": int(prev_edition[:4]), "month": int(prev_edition[5:]),
                                "filename": prev_url.split("/")[-1].split("?")[0],
                                "route": "landing_page",
                            })

    def fake_fetch(url_, dest):
        calls["fetch"] += 1
        if fail_prev and prev_edition and str(prev_edition) in str(dest):
            raise geninfo.GenInfoError("baseline unavailable")
        rows = _baseline_rows() if prev_edition and str(prev_edition) in str(dest) else _current_rows()
        _write_workbook(Path(dest), rows)
        return {"url": url_, "content_length": 910557, "last_modified": last_modified,
                "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

    monkeypatch.setattr(geninfo, "fetch_xlsx", fake_fetch)
    return calls


class TestLane:
    def test_fetch_bootstrap_diff_publish(self, tmp_path, monkeypatch):
        data_dir, docs_dir = tmp_path / "data", tmp_path / "docs" / "data"
        _install_fake_source(monkeypatch, tmp_path)

        result = run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 19))
        assert result.status == STATUS_OK
        assert result.fetched is True
        assert result.edition == "2026-07"
        assert result.row_count == 2
        counts = result.diff["counts"]
        assert counts["new_commitments"] == 1         # NEWBESS: Publicly Announced -> Committed
        assert counts["removed_units"] == 1           # CLOSED dropped out of the register
        assert result.diff["compared_to_edition"] == "2026-04"
        assert result.events_by_duid["NEWB1"][0]["type"] == "new_commitment"

        # snapshot + metadata stored; only the current raw edition kept on disk
        snapshot, meta = load_snapshot(data_dir)
        assert len(snapshot) == 2 and meta["edition"] == "2026-07"
        assert meta["previous_edition"] == "2026-04"
        raws = sorted(p.name for p in (data_dir / "geninfo").glob("*.xlsx"))
        assert raws == [f"{geninfo.FILENAME_PREFIX}2026-07.xlsx"]

        artifact = json.loads((docs_dir / geninfo.ARTIFACT_FILENAME).read_text())
        assert artifact["edition"] == "2026-07"
        assert artifact["diff"]["counts"]["new_commitments"] == 1
        assert artifact["regions"]["SA1"]["committed_mw"] == pytest.approx(250.0)

    def test_second_run_skips_download_when_rev_unchanged(self, tmp_path, monkeypatch):
        data_dir, docs_dir = tmp_path / "data", tmp_path / "docs" / "data"
        calls = _install_fake_source(monkeypatch, tmp_path)
        run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 19))
        assert calls["fetch"] == 2  # current + baseline bootstrap

        artifact = docs_dir / geninfo.ARTIFACT_FILENAME
        before = artifact.read_text()

        def _boom(url, dest):
            raise AssertionError("must not re-download an unchanged edition")

        monkeypatch.setattr(geninfo, "fetch_xlsx", _boom)
        second = run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 20))
        assert second.status == STATUS_OK
        assert second.fetched is False
        assert "rev unchanged" in second.note
        assert artifact.read_text() == before

        # forced run (explicit audit path) re-downloads regardless of rev
        forced_calls = _install_fake_source(monkeypatch, tmp_path)
        forced = run_geninfo_lane(data_dir, docs_dir, force=True, today=date(2026, 9, 20))
        assert forced.fetched is True
        assert forced_calls["fetch"] == 1

    def test_same_edition_republication_is_fetched(self, tmp_path, monkeypatch):
        data_dir, docs_dir = tmp_path / "data", tmp_path / "docs" / "data"
        _install_fake_source(monkeypatch, tmp_path)
        run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 19))
        # same edition, new rev hash => republished => refetch
        calls = _install_fake_source(
            monkeypatch, tmp_path,
            url=_JUL_URL.replace("3455851f2bc945b7ab61c5ceed272992", "deadbeef"),
        )
        result = run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 20))
        assert result.fetched is True
        assert "rev changed" in result.note
        assert calls["fetch"] == 1  # current edition only — baseline already stored

    def test_probe_route_size_and_last_modified_detection(self, tmp_path, monkeypatch):
        data_dir, docs_dir = tmp_path / "data", tmp_path / "docs" / "data"
        probe_url = geninfo.GENINFO_MEDIA_BASE + "/2026/nem-generation-information-july-2026.xlsx"
        lm = "Tue, 28 Jul 2026 00:00:00 GMT"
        _install_fake_source(monkeypatch, tmp_path, url=probe_url, last_modified=lm)
        run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 19))

        # probe route: HEAD decides via content-length + last-modified
        monkeypatch.setattr(geninfo, "probe_edition_url", lambda url: {
            "url": url, "content_length": 910557, "last_modified": lm, "content_type": "x",
        })
        skipped = run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 20))
        assert skipped.fetched is False
        assert "source unchanged" in skipped.note

        # a republication under the same edition changes size/last-modified
        calls = _install_fake_source(monkeypatch, tmp_path, url=probe_url, last_modified=lm)
        monkeypatch.setattr(geninfo, "probe_edition_url", lambda url: {
            "url": url, "content_length": 999999, "last_modified": lm, "content_type": "x",
        })
        refetched = run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 21))
        assert refetched.fetched is True
        assert "republished" in refetched.note
        assert calls["fetch"] == 1

    def test_degrades_and_retains_on_discovery_failure(self, tmp_path, monkeypatch):
        data_dir, docs_dir = tmp_path / "data", tmp_path / "docs" / "data"
        _install_fake_source(monkeypatch, tmp_path)
        run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 19))

        monkeypatch.setattr(geninfo, "discover_edition", lambda today=None: None)
        result = run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 20))
        assert result.status == STATUS_DEGRADED
        assert result.retained is True
        assert len(result.snapshot) == 2  # last-known-good kept for the attach layer

    def test_errors_on_cold_start_failure(self, tmp_path, monkeypatch):
        data_dir, docs_dir = tmp_path / "data", tmp_path / "docs" / "data"
        monkeypatch.setattr(geninfo, "discover_edition", lambda today=None: None)
        result = run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 20))
        assert result.status == STATUS_ERROR
        assert result.snapshot.empty

    def test_parse_failure_retains_previous_snapshot(self, tmp_path, monkeypatch):
        data_dir, docs_dir = tmp_path / "data", tmp_path / "docs" / "data"
        _install_fake_source(monkeypatch, tmp_path)
        run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 19))

        def _bad_fetch(url, dest):
            Path(dest).write_bytes(b"<html>not a workbook</html>")
            return {}

        monkeypatch.setattr(geninfo, "fetch_xlsx", _bad_fetch)
        result = run_geninfo_lane(data_dir, docs_dir, force=True, today=date(2026, 9, 20))
        assert result.status == STATUS_DEGRADED
        assert result.retained is True
        assert result.error and "2026-07" in result.error
        # last-known-good still stored and publishable
        snapshot, meta = load_snapshot(data_dir)
        assert len(snapshot) == 2 and meta["edition"] == "2026-07"

    def test_baseline_bootstrap_failure_publishes_baseline_only(self, tmp_path, monkeypatch):
        data_dir, docs_dir = tmp_path / "data", tmp_path / "docs" / "data"
        _install_fake_source(monkeypatch, tmp_path, fail_prev=True)
        result = run_geninfo_lane(data_dir, docs_dir, today=date(2026, 9, 19))
        assert result.status == STATUS_OK
        assert result.diff["baseline"] is True
        assert result.diff["counts"]["new_commitments"] == 0
        artifact = json.loads((docs_dir / geninfo.ARTIFACT_FILENAME).read_text())
        assert artifact["edition"] == "2026-07"


# ────────────────────────────────────────────────────────────────────────
# Run-status integration (manifest + continuity guard)
# ────────────────────────────────────────────────────────────────────────

class TestRunStatusIntegration:
    def test_gen_info_registered_as_factor_block(self):
        from src.run_status import FACTOR_LANES
        assert FACTOR_LANES["gen_info"] == "geninfo"

    def test_lane_asof_from_edition_column(self):
        lane = LaneRun(source="geninfo", block="gen_info")
        lane.frame = pd.DataFrame({"edition": ["2026-07", "2026-07"], "duid": ["A", None]})
        assert lane.asof_month() == "2026-07"

    def test_continuity_guard_protects_gen_info_blocks(self, tmp_path):
        gen_dir = tmp_path / "generators"
        gen_dir.mkdir(parents=True)
        (gen_dir / "A1.json").write_text(json.dumps({
            "duid": "A1", "gen_info": {"scope": "aemo_generation_information",
                                       "units": [{"unit_id": "1"}]},
        }))
        empty_lane = LaneRun(source="geninfo", block="gen_info", status=STATUS_DEGRADED)
        empty_lane.frame = pd.DataFrame()
        violations = check_factor_block_continuity(gen_dir, {"geninfo": empty_lane})
        assert violations and "gen_info" in violations[0]

        retained_lane = LaneRun(source="geninfo", block="gen_info", status=STATUS_DEGRADED)
        retained_lane.frame = pd.DataFrame({"duid": ["A1"], "edition": ["2026-07"]})
        assert check_factor_block_continuity(gen_dir, {"geninfo": retained_lane}) == []
